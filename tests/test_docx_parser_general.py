"""三级评分标准解析测试（用 python-docx 合成 docx）。"""
from docx import Document

from grading_agent.docx_parser import parse_standards


def _make_standards_docx(path, with_session: bool = True):
    doc = Document()
    if with_session:
        doc.add_paragraph("上午场")
        doc.add_paragraph("第一题")
        doc.add_paragraph("有定义变量(5分)")
        doc.add_paragraph("有输入语句(5分)")
        doc.add_paragraph("第二题")
        doc.add_paragraph("有循环语句(5分)")
        doc.add_paragraph("循环条件正确(5分)")
        doc.add_paragraph("下午场")
        doc.add_paragraph("第一题")
        doc.add_paragraph("有输出语句(5分)")
        doc.add_paragraph("输出内容正确(5分)")
        doc.add_paragraph("第二题")
        doc.add_paragraph("有条件语句(5分)")
        doc.add_paragraph("条件表达式正确(5分)")
    else:
        doc.add_paragraph("第一题")
        doc.add_paragraph("有定义变量(5分)")
        doc.add_paragraph("第二题")
        doc.add_paragraph("有输出语句(5分)")
    doc.save(str(path))


def test_parse_three_level_with_session(tmp_path):
    p = tmp_path / "三级评分标准.docx"
    _make_standards_docx(p, with_session=True)
    entries = parse_standards(p, level="三级")
    # 上午 2 题 + 下午 2 题 = 4 条
    assert len(entries) == 4
    assert all(e.triple.level == "三级" for e in entries)
    triples = {str(e.triple) for e in entries}
    assert triples == {"三级/上午/题1", "三级/上午/题2", "三级/下午/题1", "三级/下午/题2"}


def test_parse_three_level_no_session(tmp_path):
    p = tmp_path / "三级评分标准.docx"
    _make_standards_docx(p, with_session=False)
    entries = parse_standards(p, level="三级")
    # 未分场次：2 题复制成上下午 = 4 条
    assert len(entries) == 4
    triples = {str(e.triple) for e in entries}
    assert triples == {"三级/上午/题1", "三级/下午/题1", "三级/上午/题2", "三级/下午/题2"}


def test_parse_three_level_ten_point_dimension(tmp_path):
    """满分 10 分的维度也应正确解析。"""
    p = tmp_path / "三级评分标准.docx"
    doc = Document()
    doc.add_paragraph("第一题")
    doc.add_paragraph("综合实现(10分)")
    doc.add_paragraph("有输入语句(5分)")
    doc.add_paragraph("第二题")
    doc.add_paragraph("有输出语句(5分)")
    doc.save(str(p))
    entries = parse_standards(p, level="三级")
    # 未分场次：2 题复制成上下午 = 4 条
    assert len(entries) == 4
    # 题1 的 total_max = 10 + 5 = 15
    q1 = next(e for e in entries if e.triple.question == 1)
    assert q1.total_max == 15
    assert any(d.max_score == 10 for d in q1.dimensions)


def test_fullwidth_parentheses_dimension(tmp_path):
    """全角括号的维度也要解析（26.10.18 一级题2 回归）。

    该文档首条写作「程序框架正确（2分）」，半角正则时代被静默丢弃，
    导致 25 分题只剩 5 条、总分上限被压到 23。
    """
    p = tmp_path / "一级评分标准.docx"
    doc = Document()
    doc.add_paragraph("第一题 25分")
    doc.add_paragraph("程序框架正确（2分）")
    doc.add_paragraph("有定义变量(4分)")
    doc.add_paragraph("有输入语句(4分)")
    doc.add_paragraph("有条件语句(多if/if-else)(5分)")
    doc.add_paragraph("条件语句表达式正确(5分)")
    doc.add_paragraph("输出语句正确，且输出内容符合题意(5分)")
    doc.save(str(p))

    entries = parse_standards(p, level="一级")
    q1 = next(e for e in entries if e.triple.question == 1)
    assert len(q1.dimensions) == 6
    assert q1.total_max == 25
    assert q1.dimensions[0].name == "程序框架正确"
    assert q1.dimensions[0].max_score == 2
    # 维度名里自带的括号不应被误当分值位置
    assert any(d.name == "有条件语句(多if/if-else)" and d.max_score == 5
               for d in q1.dimensions)


def test_declared_total_mismatch_warns(tmp_path, capsys):
    """解析合计与标题声明分不符时告警，但不阻断解析。"""
    p = tmp_path / "一级评分标准.docx"
    doc = Document()
    doc.add_paragraph("第一题 25分")
    doc.add_paragraph("有定义变量(5分)")
    doc.save(str(p))

    entries = parse_standards(p, level="一级")
    err = capsys.readouterr().err
    assert "25" in err and "不符" in err
    assert entries[0].total_max == 5

