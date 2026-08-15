"""标准库通用构建 + 文件发现测试。"""
import pytest
from docx import Document

from grading_agent.standard_library import (
    StandardLibrary,
    StandardLibraryError,
    find_task_docxs,
)


def _make_task_dir(tmp_path):
    """合成一个含一级+三级评分标准、参考代码的任务文件夹。"""
    task = tmp_path / "task"
    task.mkdir()

    # 一级评分标准（不分场次，2 题）
    d = Document()
    d.add_paragraph("第一题")
    d.add_paragraph("有定义变量(5分)")
    d.add_paragraph("第二题")
    d.add_paragraph("有输出语句(5分)")
    d.save(str(task / "一级评分标准.docx"))

    # 三级评分标准（不分场次，2 题）
    d = Document()
    d.add_paragraph("第一题")
    d.add_paragraph("有输入语句(5分)")
    d.add_paragraph("第二题")
    d.add_paragraph("有循环语句(5分)")
    d.save(str(task / "三级评分标准.docx"))

    # 参考代码：上午场 + 下午场，一级 2 题 + 三级 2 题（共 8 段）
    d = Document()
    d.add_paragraph("周六上午场参考代码：")
    d.add_paragraph("一级：")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int a;"
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int b;"
    d.add_paragraph("三级：")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int c;"
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int d;"
    d.add_paragraph("周六下午场参考代码：")
    d.add_paragraph("一级：")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int e;"
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int f;"
    d.add_paragraph("三级：")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int g;"
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int h;"
    d.save(str(task / "参考代码.docx"))
    return task


def test_find_task_docxs(tmp_path):
    task = _make_task_dir(tmp_path)
    standards, ref = find_task_docxs(task)
    assert set(standards.keys()) == {"一级", "三级"}
    assert "参考代码" in ref.name


def test_standard_library_from_task_dir(tmp_path):
    task = _make_task_dir(tmp_path)
    lib = StandardLibrary.from_task_dir(task)
    triples = lib.all_triples()
    # 一级（不分场次 2 题 → 4 条）+ 三级（不分场次 2 题 → 4 条）= 8 个三元组
    assert len(triples) == 8
    # 排序：一级在前，三级在后
    levels = [t.level for t in triples]
    assert levels[:4] == ["一级", "一级", "一级", "一级"]
    assert levels[4:] == ["三级", "三级", "三级", "三级"]


def test_find_task_docxs_missing_ref(tmp_path):
    task = tmp_path / "task"
    task.mkdir()
    d = Document()
    d.add_paragraph("第一题")
    d.add_paragraph("有定义变量(5分)")
    d.save(str(task / "一级评分标准.docx"))
    with pytest.raises(FileNotFoundError):
        find_task_docxs(task)


def test_find_task_docxs_duplicate_level(tmp_path):
    task = tmp_path / "task"
    task.mkdir()
    d = Document()
    d.add_paragraph("第一题")
    d.add_paragraph("有定义变量(5分)")
    d.save(str(task / "一级评分标准.docx"))
    d.save(str(task / "一级评分标准2.docx"))
    d2 = Document()
    d2.add_paragraph("参考代码")
    d2.save(str(task / "参考代码.docx"))
    with pytest.raises(FileNotFoundError):
        find_task_docxs(task)


def test_standard_library_missing_standard_warns(tmp_path):
    """参考代码里有三级、但没有三级评分标准 → 该题跳过评分（警告而非报错）。"""
    task = tmp_path / "task"
    task.mkdir()
    d = Document()
    d.add_paragraph("第一题")
    d.add_paragraph("有定义变量(5分)")
    d.save(str(task / "一级评分标准.docx"))
    d = Document()
    d.add_paragraph("周六上午场参考代码：")
    d.add_paragraph("一级：")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int a;"
    d.add_paragraph("三级：")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int c;"
    d.add_paragraph("周六下午场参考代码：")
    d.add_paragraph("一级：")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int b;"
    d.save(str(task / "参考代码.docx"))
    lib = StandardLibrary.from_task_dir(task)
    # 一级 1 题（复制上下午）→ 2 条可评分；三级 1 题无标准 → 跳过
    assert len(lib.gradable_triples()) == 2
    assert len(lib.all_triples()) == 3
