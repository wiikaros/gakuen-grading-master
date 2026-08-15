"""三级/任意级别参考代码解析测试。"""
from docx import Document

from grading_agent.reference_codes import parse_reference_codes


def _make_ref_docx(path):
    doc = Document()
    doc.add_paragraph("周六上午场参考代码：")
    doc.add_paragraph("一级：")
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int a; cin>>a; cout<<a;"
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int b; cin>>b; cout<<b*2;"
    doc.add_paragraph("三级：")
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int c; cin>>c; cout<<c*3;"
    doc.add_paragraph("周六下午场参考代码：")
    doc.add_paragraph("三级：")
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int d; cin>>d; cout<<d*4;"
    doc.save(str(path))


def test_parse_reference_codes_three_level(tmp_path):
    p = tmp_path / "参考代码.docx"
    _make_ref_docx(p)
    refs = parse_reference_codes(p)
    keys = {str(k): v for k, v in refs.items()}
    assert set(keys.keys()) == {
        "一级/上午/题1", "一级/上午/题2",
        "三级/上午/题1", "三级/下午/题1",
    }
    assert "cout<<c*3" in keys["三级/上午/题1"]
    assert "cout<<d*4" in keys["三级/下午/题1"]


def test_parse_reference_codes_arabic_level(tmp_path):
    """阿拉伯数字级别 '3级：' 也应归一化为 '三级'。"""
    p = tmp_path / "参考代码.docx"
    doc = Document()
    doc.add_paragraph("上午场参考代码：")
    doc.add_paragraph("3级：")
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "int x;"
    doc.save(str(p))
    refs = parse_reference_codes(p)
    assert any(str(k) == "三级/上午/题1" for k in refs)
